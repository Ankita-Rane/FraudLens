"""Create and maintain the follow-up experiment Excel execution ledger."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
import json
import platform
import subprocess
import sys
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BOOK = ROOT / "artifacts" / "followup_execution_ledger.xlsx"
COMMAND_LOG_DIR = ROOT / "log" / "execution_ledger"

PLAN = [
    ("S00", "Ledger and machine readiness", "Record machine, dependencies, MPS availability, disk, source hashes", "MPS must be available before GPU stages", "pending"),
    ("S01", "Archive current follow-up artifacts", "Archive detectors, panels, attributions, feasibility and prior protocols with SHA-256 manifests", "Archive hash verification passes", "pending"),
    ("S02", "Fresh ULB EDA and detector", "Run ULB study with declared 35% test split into a new artifact directory", "Split/metrics/model/prediction manifests complete", "pending"),
    ("S03", "Fresh Sparkov EDA and detector", "Run time-aware Sparkov study into a separate artifact directory", "Split/metrics/model/prediction manifests complete", "pending"),
    ("S04", "Panel feasibility", "Exclude exposed rows and test exact 97/53 and 260/140 quotas", "Both datasets feasible; otherwise stop", "pending"),
    ("S05", "Build exact panel", "Create 150 ULB and 400 Sparkov label-separated generation/evaluation files", "550 unique cases; hashes and leakage tests pass", "pending"),
    ("S06", "TreeSHAP evidence", "Build checked attribution and diagnostic evidence for all 550 cases", "Additivity, mapping, coverage and hash checks pass", "pending"),
    ("S07", "Retrieval diagnostic re-audit", "Re-run audit against unchanged registry/corpus/model/parameters", "All scope-safe diagnostic checks pass", "pending"),
    ("S08", "Candidate protocol", "Generate and inspect candidate protocol", "550 cases, 6 conditions, repeat=1, 55 batches, 3,300 calls", "pending"),
    ("S09", "Freeze implementation and lock", "Run full tests, review changes, commit frozen inputs/code, require clean worktree, create lock", "Tests pass; clean commit; lock hashes verified", "pending"),
    ("S10", "Non-reportable smoke runs", "Run demo and Qwen/MPS smoke tests on non-panel cases", "Device, memory, persistence, pairing and schemas pass", "pending"),
    ("S11", "Final 55-batch Qwen run", "Execute batches 001-055 with immutable attempts and per-batch validation", "Each batch has 60 valid paired calls", "pending"),
    ("S12", "Aggregation and analysis", "Require exactly 3,300 calls; analyse ULB and Sparkov separately", "Completeness/hash/statistical gates pass", "pending"),
    ("S13", "Thesis evidence update", "Update tables, figures, conclusions and presentation only from accepted artifacts", "All reported values trace to accepted hashes", "pending"),
]


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def file_sha(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def style(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in range(1, sheet.max_column + 1):
            width = max(len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 100) + 1))
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 55)


def init_book(path: Path) -> None:
    if path.exists():
        raise FileExistsError(f"Refusing to overwrite execution ledger: {path}")
    workbook = Workbook()
    plan = workbook.active
    plan.title = "Plan"
    plan.append(["stage_id", "stage", "work", "completion_gate", "status", "started_utc", "completed_utc", "notes"])
    for row in PLAN:
        plan.append([*row, None, None, None])
    commands = workbook.create_sheet("Command Log")
    commands.append(["entry_id", "stage_id", "started_utc", "completed_utc", "command", "purpose", "cwd", "exit_code", "duration_seconds", "stdout_stderr_log", "log_sha256", "record_type", "notes"])
    machine = workbook.create_sheet("Machine Profile")
    machine.append(["recorded_utc", "item", "value", "status_or_use"])
    for item, value, status in [
        ("Model", "MacBook Pro Mac17,9", "observed"),
        ("Chip", "Apple M5 Pro", "observed"),
        ("CPU", "15 cores: 10 performance, 5 efficiency", "Use for pandas/scikit-learn/TF-IDF/orchestration"),
        ("GPU", "16-core integrated Apple GPU; Metal supported", "Use MPS for SBERT and Qwen after availability gate"),
        ("Unified memory", "24 GB", "Avoid concurrent model processes; monitor pressure"),
        ("PyTorch", "2.13.0", "observed"),
        ("torch.backends.mps.is_built", "true", "observed"),
        ("torch.backends.mps.is_available", "false", "BLOCKER: diagnose before GPU stages"),
        ("Free disk at initialization", "approximately 810 GiB", "sufficient"),
    ]:
        machine.append([now(), item, value, status])
    usage = workbook.create_sheet("Machine Usage Plan")
    usage.append(["workload", "preferred_device", "utilization_plan", "verification", "fallback_rule"])
    for row in [
        ("ULB/Sparkov pandas + scikit-learn", "CPU", "Use all safe estimator/joblib threads where deterministic settings permit", "Record process CPU/time/RSS", "Do not force unsupported GPU conversion"),
        ("TreeSHAP for sklearn random forests", "CPU", "Parallelize only if deterministic and memory-safe", "Record runtime/RSS and coverage", "Reduce workers on memory pressure"),
        ("TF-IDF retrieval", "CPU", "Vectorized scikit-learn execution", "Record latency/RSS", "None"),
        ("SBERT dense retrieval", "MPS", "Explicit MPS device after smoke verification", "Persist requested and observed device", "CPU only as documented non-final fallback"),
        ("Qwen generation", "MPS", "Single model process using Metal/unified memory; tune batch only by smoke evidence", "Require MPS availability and run trace", "Stop final run rather than silently use CPU"),
        ("55 experiment batches", "MPS + CPU orchestration", "Sequential validated batches; avoid competing GPU processes", "Per-batch latency/memory/device logs", "Pause on thermal/memory/error gate"),
    ]:
        usage.append(row)
    gates = workbook.create_sheet("Gates")
    gates.append(["gate_id", "stage_id", "requirement", "status", "evidence_path", "evidence_sha256", "checked_utc", "notes"])
    artifacts = workbook.create_sheet("Artifacts")
    artifacts.append(["recorded_utc", "stage_id", "artifact_path", "sha256", "bytes", "scientific_status", "relationship", "notes"])
    style(workbook)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def append_command(path: Path, values: list[object]) -> None:
    workbook = load_workbook(path)
    sheet = workbook["Command Log"]
    sheet.append(values)
    style(workbook)
    workbook.save(path)


def finish_command(path: Path, entry: str, values: dict[str, object]) -> None:
    workbook = load_workbook(path)
    sheet = workbook["Command Log"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    target = next(row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value == entry)
    for key, value in values.items():
        sheet.cell(target, headers[key], value)
    style(workbook)
    workbook.save(path)


def run_logged(args: argparse.Namespace) -> int:
    book = args.book.resolve()
    COMMAND_LOG_DIR.mkdir(parents=True, exist_ok=True)
    entry = f"CMD-{int(time.time() * 1000)}"
    started = now()
    log_path = COMMAND_LOG_DIR / f"{entry}.log"
    append_command(book, [entry, args.stage, started, None, args.command, args.purpose,
                          str(ROOT), "RUNNING", None, str(log_path.relative_to(ROOT)),
                          None, "contemporaneous", args.notes])
    before = time.monotonic()
    completed = subprocess.run(args.command, shell=True, cwd=ROOT, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    duration = time.monotonic() - before
    log_path.write_text(completed.stdout, encoding="utf-8")
    finish_command(book, entry, {
        "completed_utc": now(), "exit_code": completed.returncode,
        "duration_seconds": round(duration, 6), "log_sha256": file_sha(log_path),
    })
    sys.stdout.write(completed.stdout)
    return completed.returncode


def record_only(args: argparse.Namespace) -> None:
    entry = f"REC-{int(time.time() * 1000)}"
    append_command(args.book.resolve(), [
        entry, args.stage, args.started_utc or now(), args.completed_utc or now(),
        args.command, args.purpose, str(ROOT), args.exit_code, args.duration_seconds,
        args.log_path, args.log_sha256, args.record_type, args.notes,
    ])


def update_stage(args: argparse.Namespace) -> None:
    workbook = load_workbook(args.book.resolve())
    sheet = workbook["Plan"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    target = next(row for row in range(2, sheet.max_row + 1)
                  if sheet.cell(row, headers["stage_id"]).value == args.stage)
    sheet.cell(target, headers["status"], args.status)
    if args.started_utc:
        sheet.cell(target, headers["started_utc"], args.started_utc)
    if args.completed_utc:
        sheet.cell(target, headers["completed_utc"], args.completed_utc)
    if args.notes:
        sheet.cell(target, headers["notes"], args.notes)
    style(workbook)
    workbook.save(args.book.resolve())


def append_gate(args: argparse.Namespace) -> None:
    workbook = load_workbook(args.book.resolve())
    workbook["Gates"].append([
        args.gate_id, args.stage, args.requirement, args.status, args.evidence_path,
        args.evidence_sha256, now(), args.notes,
    ])
    style(workbook)
    workbook.save(args.book.resolve())


def append_artifact(args: argparse.Namespace) -> None:
    artifact = (ROOT / args.path).resolve() if not args.path.is_absolute() else args.path.resolve()
    if not artifact.is_file():
        raise FileNotFoundError(artifact)
    workbook = load_workbook(args.book.resolve())
    workbook["Artifacts"].append([
        now(), args.stage, str(artifact.relative_to(ROOT)), file_sha(artifact),
        artifact.stat().st_size, args.scientific_status, args.relationship, args.notes,
    ])
    style(workbook)
    workbook.save(args.book.resolve())


def append_machine(args: argparse.Namespace) -> None:
    workbook = load_workbook(args.book.resolve())
    workbook["Machine Profile"].append([now(), args.item, args.value, args.status_or_use])
    style(workbook)
    workbook.save(args.book.resolve())


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--book", type=Path, default=DEFAULT_BOOK)
    sub = parser.add_subparsers(dest="action", required=True)
    sub.add_parser("init")
    run = sub.add_parser("run")
    run.add_argument("--stage", required=True)
    run.add_argument("--purpose", required=True)
    run.add_argument("--command", required=True)
    run.add_argument("--notes", default="")
    record = sub.add_parser("record")
    record.add_argument("--stage", required=True)
    record.add_argument("--purpose", required=True)
    record.add_argument("--command", required=True)
    record.add_argument("--started-utc")
    record.add_argument("--completed-utc")
    record.add_argument("--exit-code", type=int, default=0)
    record.add_argument("--duration-seconds", type=float)
    record.add_argument("--log-path", default="")
    record.add_argument("--log-sha256", default="")
    record.add_argument("--record-type", default="retrospective")
    record.add_argument("--notes", default="")
    stage = sub.add_parser("stage")
    stage.add_argument("--stage", required=True)
    stage.add_argument("--status", required=True, choices=["pending", "in_progress", "completed", "blocked"])
    stage.add_argument("--started-utc")
    stage.add_argument("--completed-utc")
    stage.add_argument("--notes", default="")
    gate = sub.add_parser("gate")
    gate.add_argument("--gate-id", required=True)
    gate.add_argument("--stage", required=True)
    gate.add_argument("--requirement", required=True)
    gate.add_argument("--status", required=True)
    gate.add_argument("--evidence-path", default="")
    gate.add_argument("--evidence-sha256", default="")
    gate.add_argument("--notes", default="")
    artifact = sub.add_parser("artifact")
    artifact.add_argument("--stage", required=True)
    artifact.add_argument("--path", type=Path, required=True)
    artifact.add_argument("--scientific-status", required=True)
    artifact.add_argument("--relationship", required=True)
    artifact.add_argument("--notes", default="")
    machine = sub.add_parser("machine")
    machine.add_argument("--item", required=True)
    machine.add_argument("--value", required=True)
    machine.add_argument("--status-or-use", required=True)
    args = parser.parse_args()
    if args.action == "init":
        init_book(args.book.resolve())
        print(json.dumps({"book": str(args.book.resolve()), "sha256": file_sha(args.book.resolve())}, indent=2))
    elif args.action == "run":
        raise SystemExit(run_logged(args))
    elif args.action == "record":
        record_only(args)
    elif args.action == "stage":
        update_stage(args)
    elif args.action == "gate":
        append_gate(args)
    elif args.action == "artifact":
        append_artifact(args)
    else:
        append_machine(args)


if __name__ == "__main__":
    main()
